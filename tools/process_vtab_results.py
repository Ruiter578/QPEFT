import os
import argparse
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def parse_experiment_folder(root_path):
    """
    Parses the experiment folder to extract dataset names and accuracies.
    
    Args:
        root_path (str): The path to the root folder containing dataset subfolders.
        
    Returns:
        dict: A dictionary mapping dataset names to their accuracies.
    """
    results = {}
    if not os.path.isdir(root_path):
        print(f"Error: Directory not found at '{root_path}'")
        return None

    # List all items in the root path
    for dataset_name in os.listdir(root_path):
        dataset_path = os.path.join(root_path, dataset_name)
        
        # Ensure it's a directory
        if os.path.isdir(dataset_path):
            # Find the inner folder with the accuracy score
            try:
                inner_folders = [d for d in os.listdir(dataset_path) if os.path.isdir(os.path.join(dataset_path, d))]
                if not inner_folders:
                    print(f"Warning: No result folder found in '{dataset_path}'")
                    continue
                
                # Assume the first folder found is the correct one
                result_folder_name = inner_folders[0]
                
                # Use regex to find '_acc=<number>'
                match = re.search(r'_acc=([\d\.]+)', result_folder_name)
                if match:
                    accuracy = float(match.group(1))
                    results[dataset_name] = accuracy
                else:
                    print(f"Warning: Could not parse accuracy from '{result_folder_name}' in '{dataset_path}'")

            except Exception as e:
                print(f"Error processing directory {dataset_path}: {e}")
                
    return results

def write_to_excel(results, output_file):
    """
    Writes the parsed results to an Excel file with specific formatting.
    
    Args:
        results (dict): Dictionary of dataset accuracies.
        output_file (str): Path for the output .xlsx file.
    """
    # Define the structure and order from the target image
    # Note: Using the exact names from your folder structure
    header_map = {
        'cifar': 'Cifar100', 'caltech101': 'Caltech101', 'dtd': 'DTD', 
        'oxford_flowers102': 'Flowers102', 'oxford_iiit_pet': 'Pets', 'svhn': 'SVHN', 'sun397': 'Sun397',
        'patch_camelyon': 'Camelyon', 'eurosat': 'EuroSAT', 'resisc45': 'Resisc45', 'diabetic_retinopathy': 'Retinopathy',
        'clevr_count': 'Clevr-Count', 'clevr_dist': 'Clevr-Dist', 'dmlab': 'DMLAB', 'kitti': 'KITTI-Dist',
        'dsprites_loc': 'd Sprites-Loc', 'dsprites_ori': 'd Spr-Ori', 'smallnorb_azi': 'sNORB-Azim', 'smallnorb_ele': 'sNORB-Ele'
    }

    natural_group = ['cifar', 'caltech101', 'dtd', 'oxford_flowers102', 'oxford_iiit_pet', 'svhn', 'sun397']
    specialized_group = ['patch_camelyon', 'eurosat', 'resisc45', 'diabetic_retinopathy']
    structured_group = ['clevr_count', 'clevr_dist', 'dmlab', 'kitti', 'dsprites_loc', 'dsprites_ori', 'smallnorb_azi', 'smallnorb_ele']

    wb = Workbook()
    ws = wb.active
    ws.title = "VTAB Results"

    # --- Define Styles ---
    red_font = Font(color="FF0000", bold=True)
    center_align = Alignment(horizontal='center', vertical='center')

    # --- Create Headers ---
    headers = []
    col_idx = 1

    # Natural Group
    start_col = col_idx
    ws.cell(row=1, column=start_col, value="Natural").alignment = center_align
    for key in natural_group:
        headers.append(key)
        ws.cell(row=2, column=col_idx, value=header_map.get(key, key))
        col_idx += 1
    headers.append('natural_avg')
    ws.cell(row=2, column=col_idx, value="Group Avg").font = red_font
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col_idx)
    col_idx += 1

    # Specialized Group
    start_col = col_idx
    ws.cell(row=1, column=start_col, value="Specialized").alignment = center_align
    for key in specialized_group:
        headers.append(key)
        ws.cell(row=2, column=col_idx, value=header_map.get(key, key))
        col_idx += 1
    headers.append('specialized_avg')
    ws.cell(row=2, column=col_idx, value="Group Avg").font = red_font
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col_idx)
    col_idx += 1

    # Structured Group
    start_col = col_idx
    ws.cell(row=1, column=start_col, value="Structured").alignment = center_align
    for key in structured_group:
        headers.append(key)
        ws.cell(row=2, column=col_idx, value=header_map.get(key, key))
        col_idx += 1
    headers.append('structured_avg')
    ws.cell(row=2, column=col_idx, value="Group Avg").font = red_font
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col_idx)
    col_idx += 1
    
    # Group Averages Mean
    headers.append('group_avg_mean')
    ws.cell(row=2, column=col_idx, value="Group Avg Mean").font = red_font
    col_idx += 1
    
    # Total Average
    headers.append('total_avg')
    ws.cell(row=2, column=col_idx, value="Tot. Avg").font = red_font
    
    # --- Calculate Averages ---
    all_accuracies = [v for v in results.values() if v is not None]

    def get_avg(group):
        vals = [results.get(k) for k in group if results.get(k) is not None]
        return sum(vals) / len(vals) if vals else 0

    natural_avg = get_avg(natural_group)
    specialized_avg = get_avg(specialized_group)
    structured_avg = get_avg(structured_group)
    
    calcs = {
        'natural_avg': natural_avg,
        'specialized_avg': specialized_avg,
        'structured_avg': structured_avg,
        'group_avg_mean': (natural_avg + specialized_avg + structured_avg) / 3,
        'total_avg': sum(all_accuracies) / len(all_accuracies) if all_accuracies else 0
    }

    # --- Write Data Row ---
    data_row = 3
    for i, header_key in enumerate(headers, 1):
        cell = ws.cell(row=data_row, column=i)
        if 'avg' in header_key:
            value = calcs.get(header_key)
            cell.value = f"{value:.2f}" if value is not None else "N/A"
            cell.font = red_font
        else:
            value = results.get(header_key)
            cell.value = value if value is not None else "N/A"

    # Adjust column widths for better readability
    for i in range(1, col_idx + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12

    # --- Save the file ---
    try:
        wb.save(output_file)
        print(f"Successfully created Excel report at: {output_file}")
    except Exception as e:
        print(f"Error saving Excel file: {e}")

def main():
    """Main function to run the script."""
    parser = argparse.ArgumentParser(
        description="Extract experiment results from a folder structure and generate a formatted Excel report.",
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument(
        "folder_path",
        help="Path to the root experiment folder (e.g., 'T=vtab_Q=None_L=lora_2_proj-fc2')."
    )
    # parser.add_argument(
    #     "-o", "--output",
    #     default="vtab_results.xlsx",
    #     help="Name of the output Excel file (default: vtab_results.xlsx)."
    # )
    
    args = parser.parse_args()
    
    results_data = parse_experiment_folder(args.folder_path)
    
    if results_data:
        output_file = os.path.join(args.folder_path, "vtab_results.xlsx")
        write_to_excel(results_data, output_file)

if __name__ == "__main__":
    main()