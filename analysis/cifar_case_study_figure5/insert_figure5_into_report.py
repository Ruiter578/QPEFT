from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DOCX = ROOT / "Final_Report_QPEFT_revised_v6检查版.docx"
DEFAULT_FIGURE = ROOT / "analysis/cifar_case_study_figure5/figures/figure5_cifar_3bit_qpeft_val_acc1_trajectory.png"
DEFAULT_OUTPUT = ROOT / "Final_Report_QPEFT_revised_v6检查版_with_Figure5.docx"


def insert_paragraph_before(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    return Paragraph(new_p, paragraph._parent)


def main() -> None:
    doc = Document(str(DEFAULT_DOCX))
    target = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Figure 5."):
            target = paragraph
            break
    if target is None:
        raise RuntimeError("Could not find a paragraph starting with 'Figure 5.'")

    figure_paragraph = insert_paragraph_before(target)
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = figure_paragraph.add_run()
    run.add_picture(str(DEFAULT_FIGURE), width=Inches(6.2))

    doc.save(str(DEFAULT_OUTPUT))
    print(f"Inserted Figure 5 into: {DEFAULT_OUTPUT}")


if __name__ == "__main__":
    main()
